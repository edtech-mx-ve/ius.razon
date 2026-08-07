from __future__ import annotations

import json
from datetime import date
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from ius_razon.config import AppConfig
from ius_razon.domain.argumentation_models import (
    ArgumentPosition,
    ArgumentRelationCreate,
    ArgumentRelationType,
    ArgumentScenarioCreate,
    ArgumentStatus,
    LegalArgumentCreate,
    ScenarioStatus,
)
from ius_razon.domain.enums import (
    CaseStatus,
    ConfidentialityLevel,
    EvidenceEvaluationStatus,
    EvidenceType,
    FactStatus,
    LegalIssueStatus,
    LegalSourceType,
    NormHierarchy,
    PartyRole,
    PartyType,
    SourceOrientation,
)
from ius_razon.domain.models import (
    CaseCreate,
    EvidenceCreate,
    FactCreate,
    IssueSourceLinkCreate,
    LegalIssueCreate,
    NormCreate,
    PartyCreate,
)
from ius_razon.domain.reasoning_models import (
    AssertionValue,
    ConditionRole,
    ReasoningAssertionCreate,
    ReasoningRuleCreate,
    RuleConditionCreate,
    RuleKind,
)
from ius_razon.domain.report_models import IntegralReportRequest
from ius_razon.persistence.argumentation_repository import (
    ArgumentationRepository,
)
from ius_razon.persistence.mutation_backup import SQLiteMutationBackup
from ius_razon.persistence.reasoning_repository import ReasoningRepository
from ius_razon.persistence.sqlite_repository import SQLiteRepository
from ius_razon.services.argumentation_service import ArgumentationService
from ius_razon.services.case_service import CaseService
from ius_razon.services.legal_report_service import LegalReportService
from ius_razon.services.reasoning_service import ReasoningService


def build_services(
    tmp_path: Path,
) -> tuple[
    CaseService,
    ReasoningService,
    ArgumentationService,
    LegalReportService,
]:
    data_dir = tmp_path / "data"
    config = AppConfig(
        project_root=tmp_path,
        data_dir=data_dir,
        db_path=data_dir / "test.db",
        upload_dir=data_dir / "uploads",
        log_dir=data_dir / "logs",
        max_upload_bytes=1024 * 1024,
        log_level="INFO",
    )
    config.upload_dir.mkdir(parents=True)
    config.log_dir.mkdir(parents=True)
    case_repository = SQLiteRepository(config.db_path)
    case_repository.initialize()
    reasoning_repository = ReasoningRepository(config.db_path)
    reasoning_repository.initialize()
    argumentation_repository = ArgumentationRepository(config.db_path)
    argumentation_repository.initialize()
    case_service = CaseService(
        case_repository,
        config,
        mutation_backup=SQLiteMutationBackup(
            config.db_path,
            config.data_dir / "backups",
            keep=20,
        ),
    )
    reasoning_service = ReasoningService(
        reasoning_repository,
        mutation_backup=SQLiteMutationBackup(
            reasoning_repository.db_path,
            data_dir / "backups",
            keep=20,
        ),
    )
    argumentation_service = ArgumentationService(
        argumentation_repository,
        mutation_backup=SQLiteMutationBackup(
            argumentation_repository.db_path,
            data_dir / "backups",
            keep=20,
        ),
    )
    return (
        case_service,
        reasoning_service,
        argumentation_service,
        LegalReportService(
            case_service,
            reasoning_service,
            argumentation_service,
        ),
    )


def create_integral_case(
    tmp_path: Path,
) -> tuple[
    LegalReportService,
    IntegralReportRequest,
    str,
    str,
    str,
]:
    case_service, reasoning, argumentation, report_service = build_services(
        tmp_path
    )
    case = case_service.create_case(
        CaseCreate(
            title="Incumplimiento de entrega",
            description=(
                "Expediente sintético para validar el informe jurídico integral."
            ),
            matter="Mercantil",
            jurisdiction="México",
            opened_on=date.today(),
            status=CaseStatus.OPEN,
            objective="Revisar la trazabilidad del incumplimiento.",
            user_role="Analista",
            confidentiality=ConfidentialityLevel.PUBLIC_DEMO,
        )
    )
    buyer = case_service.add_party(
        PartyCreate(
            case_id=case.id,
            name_alias="Compradora A",
            party_type=PartyType.LEGAL_ENTITY,
            legal_role=PartyRole.CLAIMANT,
            claim="Solicita reconocer el incumplimiento.",
        )
    )
    fact = case_service.add_fact(
        FactCreate(
            case_id=case.id,
            description="El plazo venció sin entrega registrada.",
            actor_party_id=buyer.id,
            action="Pagó",
            object_text="Equipo",
            status=FactStatus.SUPPORTED,
            controversy_level=2,
        )
    )
    evidence = case_service.add_evidence(
        EvidenceCreate(
            case_id=case.id,
            evidence_type=EvidenceType.RECEIPT,
            description="Comprobante de pago total.",
            evaluation_status=EvidenceEvaluationStatus.SUPPORTED,
        ),
        original_file_name=None,
        file_content=None,
    )
    case_service.link_fact_evidence(
        case_id=case.id,
        fact_id=fact.id,
        evidence_id=evidence.id,
        purpose="Acreditar el pago registrado.",
    )
    issue = case_service.add_legal_issue(
        LegalIssueCreate(
            case_id=case.id,
            title="Incumplimiento de entrega",
            question="¿Existió incumplimiento de la obligación de entrega?",
            description="Problema jurídico sintético.",
            status=LegalIssueStatus.UNDER_ANALYSIS,
        )
    )
    norm = case_service.add_norm(
        NormCreate(
            case_id=case.id,
            jurisdiction="México",
            matter="Mercantil",
            instrument="Contrato de compraventa",
            article="Cláusula 4",
            text="La parte vendedora debe entregar el equipo en la fecha pactada.",
            hierarchy=NormHierarchy.CONTRACTUAL,
        ),
        original_file_name=None,
        file_content=None,
    )
    case_service.link_issue_source(
        IssueSourceLinkCreate(
            case_id=case.id,
            issue_id=issue.id,
            source_type=LegalSourceType.NORM,
            source_id=norm.id,
            orientation=SourceOrientation.SUPPORTS,
            applicability="Define la obligación de entrega.",
        )
    )
    assertion = reasoning.add_assertion(
        ReasoningAssertionCreate(
            case_id=case.id,
            issue_id=issue.id,
            predicate_key="entrega_incumplida",
            statement="La entrega no se realizó en el plazo.",
            value=AssertionValue.TRUE,
            basis="El hecho y la prueba sostienen inicialmente la valoración.",
            support_codes=[fact.code, evidence.code],
        )
    )
    reasoning.add_rule(
        ReasoningRuleCreate(
            case_id=case.id,
            issue_id=issue.id,
            name="Regla de incumplimiento",
            kind=RuleKind.DEFAULT,
            conclusion_key="incumplimiento_entrega",
            conclusion_statement="Existe soporte provisional de incumplimiento.",
            conclusion_value=AssertionValue.TRUE,
            priority=100,
            active=True,
            legal_basis_codes=[norm.code],
            explanation="La falta de entrega permite derivar incumplimiento.",
            conditions=[
                RuleConditionCreate(
                    role=ConditionRole.PREREQUISITE,
                    predicate_key="entrega_incumplida",
                    expected_value=AssertionValue.TRUE,
                )
            ],
        )
    )
    run = reasoning.run(case_id=case.id, issue_id=issue.id)
    favorable = argumentation.create_from_conclusion(
        conclusion_id=run.conclusions[0].id,
        position=ArgumentPosition.FAVORABLE,
        title="Argumento favorable principal",
        claim="Existe soporte provisional para reconocer el incumplimiento.",
        reasoning="La inferencia conecta el hecho, la prueba y la norma.",
        status=ArgumentStatus.SUPPORTED,
    )
    adverse = argumentation.add_argument(
        LegalArgumentCreate(
            case_id=case.id,
            issue_id=issue.id,
            title="Objeción adversa",
            position=ArgumentPosition.ADVERSE,
            thesis_key="incumplimiento_entrega",
            thesis_statement="No existe soporte suficiente.",
            thesis_value=AssertionValue.FALSE,
            claim="La objeción cuestiona la suficiencia de la inferencia.",
            reasoning="La autenticidad y aplicabilidad requieren revisión.",
            status=ArgumentStatus.OBJECTED,
            fact_codes=[fact.code],
            evidence_codes=[evidence.code],
            source_codes=[norm.code],
            rule_codes=[],
            assertion_codes=[assertion.code],
        )
    )
    reply = argumentation.add_argument(
        LegalArgumentCreate(
            case_id=case.id,
            issue_id=issue.id,
            title="Réplica a la objeción",
            position=ArgumentPosition.FAVORABLE,
            thesis_key="objecion_resuelta",
            thesis_statement="La objeción fue respondida estructuralmente.",
            thesis_value=AssertionValue.TRUE,
            claim="La réplica conserva la revisión humana como límite.",
            reasoning="La réplica explica el alcance provisional del resultado.",
            status=ArgumentStatus.ANSWERED,
            fact_codes=[fact.code],
            evidence_codes=[evidence.code],
            source_codes=[norm.code],
            rule_codes=["R-001"],
            assertion_codes=[assertion.code],
        )
    )
    argumentation.add_relation(
        ArgumentRelationCreate(
            case_id=case.id,
            issue_id=issue.id,
            source_argument_id=adverse.id,
            target_argument_id=favorable.id,
            relation_type=ArgumentRelationType.ATTACKS,
            rationale="La objeción ataca la suficiencia del argumento favorable.",
        )
    )
    argumentation.add_relation(
        ArgumentRelationCreate(
            case_id=case.id,
            issue_id=issue.id,
            source_argument_id=reply.id,
            target_argument_id=adverse.id,
            relation_type=ArgumentRelationType.REPLIES,
            rationale="La réplica responde a la objeción adversa.",
        )
    )
    base = argumentation.add_scenario(
        ArgumentScenarioCreate(
            case_id=case.id,
            issue_id=issue.id,
            name="Escenario sin réplica",
            description="Incluye la tesis y la objeción, sin respuesta.",
            status=ScenarioStatus.ACTIVE,
            argument_ids=[favorable.id, adverse.id],
            assumptions=["No se incorpora la réplica."],
        )
    )
    compared = argumentation.add_scenario(
        ArgumentScenarioCreate(
            case_id=case.id,
            issue_id=issue.id,
            name="Escenario con réplica",
            description="Incluye la tesis, la objeción y la respuesta.",
            status=ScenarioStatus.ACTIVE,
            argument_ids=[favorable.id, adverse.id, reply.id],
            assumptions=["La réplica se considera estructuralmente suficiente."],
        )
    )
    request = IntegralReportRequest(
        case_id=case.id,
        issue_id=issue.id,
        reasoning_run_id=run.run.id,
        title="Informe jurídico integral de incumplimiento",
        purpose="Integrar datos, inferencia, argumentos y escenarios.",
        base_scenario_id=base.id,
        compared_scenario_id=compared.id,
        analyst_conclusions=[
            "La conclusión permanece provisional y exige revisión humana."
        ],
        include_full_arguments=True,
        include_source_details=True,
        include_traceability_appendix=True,
    )
    return report_service, request, case.id, issue.id, run.run.id


def test_integral_report_collects_all_layers(tmp_path: Path) -> None:
    service, request, _, _, _ = create_integral_case(tmp_path)

    report = service.build_report(request)

    assert report.report_version == "4.2.1"
    assert len(report.parties) == 1
    assert len(report.facts) == 1
    assert len(report.evidence) == 1
    assert len(report.norms) == 1
    assert len(report.reasoning.conclusions) == 1
    assert len(report.argumentation.arguments) == 3
    assert report.scenario_comparison is not None
    assert report.traceability_rows[0].argument_code == "ARG-001"


def test_scenario_narrative_explains_resolved_objection(
    tmp_path: Path,
) -> None:
    service, request, _, _, _ = create_integral_case(tmp_path)

    report = service.build_report(request)

    assert report.scenario_narrative is not None
    assert "ARG-003" in " ".join(report.scenario_narrative.statements)
    assert any(
        "disminuyen en 1" in item
        for item in report.scenario_narrative.statements
    )


def test_report_hash_and_json_are_reproducible(tmp_path: Path) -> None:
    service, request, _, _, _ = create_integral_case(tmp_path)

    first = service.build_report(request)
    second = service.build_report(request)

    assert first.input_hash == second.input_hash
    assert service.export_json(request) == service.export_json(request)
    parsed = json.loads(service.export_json(request))
    assert parsed["input_hash"] == first.input_hash


def test_markdown_contains_integral_sections(tmp_path: Path) -> None:
    service, request, _, _, _ = create_integral_case(tmp_path)

    markdown = service.export_markdown(request)

    assert "## 7. Inferencia y trazabilidad" in markdown
    assert "## 8. Argumentación jurídica" in markdown
    assert "## 9. Escenarios alternativos" in markdown
    assert "Anexo A. Matriz de trazabilidad" in markdown
    assert "ARG-003" in markdown


def test_docx_is_valid_and_contains_expected_headings(tmp_path: Path) -> None:
    service, request, _, _, _ = create_integral_case(tmp_path)

    payload = service.export_docx(request)
    document = Document(BytesIO(payload))
    headings = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith("Heading")
    ]

    assert payload.startswith(b"PK")
    assert "7. Inferencia y trazabilidad" in headings
    assert "8. Argumentación jurídica" in headings
    assert "Anexo A. Matriz de trazabilidad" in headings
    assert len(document.tables) >= 4


def test_custom_summary_and_recommendations_are_preserved(
    tmp_path: Path,
) -> None:
    service, request, _, _, _ = create_integral_case(tmp_path)
    custom = request.model_copy(
        update={
            "executive_summary": "Resumen ejecutivo registrado por el analista.",
            "recommendations": ["Verificar la vigencia de la norma contractual."],
        }
    )

    report = service.build_report(custom)

    assert report.executive_summary.startswith("Resumen ejecutivo")
    assert report.recommendations == [
        "Verificar la vigencia de la norma contractual."
    ]


def test_run_from_other_issue_is_rejected(tmp_path: Path) -> None:
    service, request, case_id, _, run_id = create_integral_case(tmp_path)
    wrong = request.model_copy(
        update={
            "case_id": case_id,
            "issue_id": "problema-inexistente",
            "reasoning_run_id": run_id,
        }
    )

    with pytest.raises(ValueError, match="problema jurídico"):
        service.build_report(wrong)


def test_equal_scenarios_are_rejected_by_request(tmp_path: Path) -> None:
    _, request, _, _, _ = create_integral_case(tmp_path)

    with pytest.raises(ValueError, match="distintos"):
        IntegralReportRequest(
            **request.model_dump(
                exclude={"compared_scenario_id"},
            ),
            compared_scenario_id=request.base_scenario_id,
        )


def test_manual_arguments_are_described_without_false_traceability_gap(
    tmp_path: Path,
) -> None:
    service, request, _, _, _ = create_integral_case(tmp_path)

    report = service.build_report(request)
    manual_codes = [
        argument.code
        for argument in report.argumentation.arguments
        if argument.conclusion_id is None
    ]
    support_by_code = {
        item.argument_code: item
        for item in report.argumentation.support_assessments
    }

    assert manual_codes == ["ARG-002", "ARG-003"]
    assert all(
        "conclusión inferida" not in item
        for item in report.argumentation.missing_information
    )
    assert all(support_by_code[code].score == 5 for code in manual_codes)
    assert any(
        "argumentos manuales sin conclusiones inferidas asociadas" in item
        for item in report.findings
    )
    assert all(
        "faltan conclusión inferida" not in item
        for item in report.limitations
    )


def test_findings_apply_spanish_number_agreement(tmp_path: Path) -> None:
    service, request, _, _, _ = create_integral_case(tmp_path)

    report = service.build_report(request)
    findings = " ".join(report.findings)

    assert "1 argumento favorable y 1 argumento adverso" in findings
    assert "1 objeción pendiente" in findings
    assert "1 adversos" not in findings


def test_markdown_identifies_manual_argument_origin(tmp_path: Path) -> None:
    service, request, _, _, _ = create_integral_case(tmp_path)

    markdown = service.export_markdown(request)

    assert "- Origen: Manual" in markdown
    assert "No aplica (argumento manual)" in markdown
    assert "faltan conclusión inferida" not in markdown


def test_docx_avoids_duplicate_sentence_punctuation(tmp_path: Path) -> None:
    service, request, _, _, _ = create_integral_case(tmp_path)

    payload = service.export_docx(request)
    document = Document(BytesIO(payload))
    document_text = "\n".join(
        paragraph.text for paragraph in document.paragraphs
    )

    assert "Acreditar el pago registrado.." not in document_text
    assert "Acreditar el pago registrado." in document_text
    assert "Origen: Manual" in document_text
