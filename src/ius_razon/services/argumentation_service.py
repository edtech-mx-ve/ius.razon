from __future__ import annotations

import hashlib
import io
import json
import logging

from docx import Document

from ius_razon.domain.argumentation_models import (
    ArgumentationDossier,
    ArgumentGraphEdge,
    ArgumentGraphNode,
    ArgumentGraphSnapshot,
    ArgumentPosition,
    ArgumentRelationCreate,
    ArgumentRelationRecord,
    ArgumentRelationType,
    ArgumentScenarioCreate,
    ArgumentScenarioRecord,
    ArgumentScenarioUpdate,
    ArgumentStatus,
    ArgumentSupportAssessment,
    ArgumentSupportLevel,
    LegalArgumentCreate,
    LegalArgumentRecord,
    LegalArgumentUpdate,
    ScenarioComparison,
)
from ius_razon.domain.reasoning_models import AssertionValue
from ius_razon.persistence.argumentation_repository_protocol import (
    ArgumentationRepositoryProtocol,
)
from ius_razon.persistence.mutation_backup import MutationBackup

LOGGER = logging.getLogger(__name__)


def assess_argument_support(
    argument: LegalArgumentRecord,
) -> ArgumentSupportAssessment:
    """Evalúa completitud de trazabilidad sin estimar probabilidad jurídica."""

    thesis_origin = (
        ("conclusión inferida", True)
        if argument.conclusion_id is not None
        else ("tesis manual explícita", bool(argument.thesis_statement))
    )
    components = (
        ("hechos", bool(argument.fact_codes)),
        ("pruebas", bool(argument.evidence_codes)),
        ("fuentes jurídicas", bool(argument.source_codes)),
        (
            "reglas o premisas",
            bool(argument.rule_codes or argument.assertion_codes),
        ),
        thesis_origin,
    )
    score = sum(is_present for _, is_present in components)
    if score >= 4:
        level = ArgumentSupportLevel.HIGH
    elif score == 3:
        level = ArgumentSupportLevel.MEDIUM
    elif score == 2:
        level = ArgumentSupportLevel.LOW
    else:
        level = ArgumentSupportLevel.INSUFFICIENT
    return ArgumentSupportAssessment(
        argument_code=argument.code,
        level=level,
        score=score,
        missing_components=[
            label for label, is_present in components if not is_present
        ],
    )


def _escape_dot(value: str) -> str:
    """Escapa texto para una etiqueta DOT sin ejecutar contenido externo."""

    return (
        value.replace("\\", "\\\\")
        .replace('"', '\\"')
        .replace("\n", " ")
        .replace("\r", " ")
    )


def _connected_component_count(
    node_ids: set[str],
    relations: list[ArgumentRelationRecord],
) -> int:
    """Cuenta componentes del grafo como red no dirigida."""

    if not node_ids:
        return 0
    adjacency: dict[str, set[str]] = {
        node_id: set() for node_id in node_ids
    }
    for relation in relations:
        source = relation.source_argument_id
        target = relation.target_argument_id
        if source in adjacency and target in adjacency:
            adjacency[source].add(target)
            adjacency[target].add(source)
    remaining = set(node_ids)
    component_count = 0
    while remaining:
        component_count += 1
        pending = [remaining.pop()]
        while pending:
            current = pending.pop()
            neighbors = adjacency[current] & remaining
            remaining.difference_update(neighbors)
            pending.extend(neighbors)
    return component_count


class ArgumentationService:
    """Orquesta validación, relaciones y reportes argumentales."""

    def __init__(
        self,
        repository: ArgumentationRepositoryProtocol,
        *,
        mutation_backup: MutationBackup,
    ) -> None:
        self._repository = repository
        self._mutation_backup = mutation_backup

    def list_conclusion_contexts(
        self,
        case_id: str,
        issue_id: str,
    ) -> list[dict[str, str | list[str]]]:
        """Lista conclusiones disponibles para construir argumentos."""

        self._repository.validate_case_issue(case_id, issue_id)
        return self._repository.list_conclusion_contexts(case_id, issue_id)

    def add_argument(
        self,
        payload: LegalArgumentCreate,
    ) -> LegalArgumentRecord:
        """Valida y registra un argumento manual."""

        self._repository.validate_case_issue(payload.case_id, payload.issue_id)
        self._validate_argument_codes(payload.case_id, payload.issue_id, payload)
        self._validate_conclusion_link(payload)
        record = self._repository.add_argument(payload)
        LOGGER.info(
            "Argumento creado. case_id=%s argument_id=%s code=%s",
            payload.case_id,
            record.id,
            record.code,
        )
        return record

    def create_from_conclusion(
        self,
        *,
        conclusion_id: str,
        position: ArgumentPosition,
        title: str,
        claim: str,
        reasoning: str,
        status: ArgumentStatus = ArgumentStatus.DRAFT,
    ) -> LegalArgumentRecord:
        """Construye un argumento trazable desde una conclusión persistida."""

        context = self._repository.conclusion_context(conclusion_id)
        case_id = str(context["case_id"])
        issue_id = str(context["issue_id"])
        source_codes = [str(item) for item in context["source_codes"]]
        assertion_codes = [str(item) for item in context["assertion_codes"]]
        rule_codes = [str(item) for item in context["rule_codes"]]
        payload = LegalArgumentCreate(
            case_id=case_id,
            issue_id=issue_id,
            title=title,
            position=position,
            thesis_key=str(context["predicate_key"]),
            thesis_statement=str(context["statement"]),
            thesis_value=AssertionValue(str(context["value"])),
            claim=claim,
            reasoning=reasoning,
            status=status,
            conclusion_id=conclusion_id,
            fact_codes=[
                code for code in source_codes if code.startswith("H-")
            ],
            evidence_codes=[
                code for code in source_codes if code.startswith("P-")
            ],
            source_codes=[
                code
                for code in source_codes
                if code.startswith(("N-", "J-", "D-"))
            ],
            rule_codes=rule_codes,
            assertion_codes=assertion_codes,
        )
        return self.add_argument(payload)

    def update_argument(
        self,
        argument_id: str,
        payload: LegalArgumentUpdate,
    ) -> LegalArgumentRecord:
        """Actualiza con respaldo previo y preserva el código."""

        current = self._repository.get_argument(argument_id)
        self._validate_argument_codes(current.case_id, current.issue_id, payload)
        if payload.conclusion_id is not None:
            context = self._repository.conclusion_context(payload.conclusion_id)
            if (
                str(context["case_id"]) != current.case_id
                or str(context["issue_id"]) != current.issue_id
            ):
                raise ValueError(
                    "La conclusión vinculada pertenece a otro problema jurídico."
                )
        self._backup_before_mutation("actualizar argumento")
        return self._repository.update_argument(argument_id, payload)

    def delete_argument(
        self,
        argument_id: str,
        *,
        confirmation: str,
    ) -> None:
        """Elimina con respaldo y confirmación exacta."""

        self._backup_before_mutation("eliminar argumento")
        self._repository.delete_argument(
            argument_id,
            confirmation=confirmation,
        )

    def list_arguments(
        self,
        case_id: str,
        issue_id: str | None = None,
        *,
        include_discarded: bool = False,
    ) -> list[LegalArgumentRecord]:
        return self._repository.list_arguments(
            case_id,
            issue_id,
            include_discarded=include_discarded,
        )

    def add_relation(
        self,
        payload: ArgumentRelationCreate,
    ) -> ArgumentRelationRecord:
        """Registra apoyo, ataque o réplica entre argumentos."""

        self._repository.validate_case_issue(payload.case_id, payload.issue_id)
        self._backup_before_mutation("crear relación argumental")
        return self._repository.add_relation(payload)

    def delete_relation(
        self,
        relation_id: str,
        *,
        confirmation: str,
    ) -> None:
        """Elimina una relación con respaldo previo."""

        self._backup_before_mutation("eliminar relación argumental")
        self._repository.delete_relation(
            relation_id,
            confirmation=confirmation,
        )

    def list_relations(
        self,
        case_id: str,
        issue_id: str | None = None,
    ) -> list[ArgumentRelationRecord]:
        return self._repository.list_relations(case_id, issue_id)

    def add_scenario(
        self,
        payload: ArgumentScenarioCreate,
    ) -> ArgumentScenarioRecord:
        """Crea un escenario tras validar sus argumentos."""

        self._repository.validate_case_issue(payload.case_id, payload.issue_id)
        self._validate_scenario_arguments(
            payload.case_id,
            payload.issue_id,
            payload.argument_ids,
        )
        self._backup_before_mutation("crear escenario argumental")
        return self._repository.add_scenario(payload)

    def update_scenario(
        self,
        scenario_id: str,
        payload: ArgumentScenarioUpdate,
    ) -> ArgumentScenarioRecord:
        """Actualiza un escenario conservando su código."""

        current = self._repository.get_scenario(scenario_id)
        self._validate_scenario_arguments(
            current.case_id,
            current.issue_id,
            payload.argument_ids,
        )
        self._backup_before_mutation("actualizar escenario argumental")
        return self._repository.update_scenario(scenario_id, payload)

    def delete_scenario(
        self,
        scenario_id: str,
        *,
        confirmation: str,
    ) -> None:
        """Elimina un escenario con respaldo y confirmación exacta."""

        self._backup_before_mutation("eliminar escenario argumental")
        self._repository.delete_scenario(
            scenario_id,
            confirmation=confirmation,
        )

    def list_scenarios(
        self,
        case_id: str,
        issue_id: str | None = None,
        *,
        include_archived: bool = False,
    ) -> list[ArgumentScenarioRecord]:
        """Lista escenarios del expediente."""

        return self._repository.list_scenarios(
            case_id,
            issue_id,
            include_archived=include_archived,
        )

    def build_graph(
        self,
        case_id: str,
        issue_id: str,
        scenario_id: str | None = None,
    ) -> ArgumentGraphSnapshot:
        """Construye un grafo completo o restringido a un escenario."""

        self._repository.validate_case_issue(case_id, issue_id)
        all_arguments = self.list_arguments(case_id, issue_id)
        all_relations = self.list_relations(case_id, issue_id)
        scenario: ArgumentScenarioRecord | None = None
        selected_ids: set[str]
        scenario_code: str | None
        scenario_name: str
        assumptions: list[str]
        if scenario_id is None:
            selected_ids = {argument.id for argument in all_arguments}
            scenario_code = None
            scenario_name = "Vista completa"
            assumptions = []
        else:
            scenario = self._repository.get_scenario(scenario_id)
            if scenario.case_id != case_id or scenario.issue_id != issue_id:
                raise ValueError(
                    "El escenario pertenece a otro problema jurídico."
                )
            selected_ids = set(scenario.argument_ids)
            scenario_code = scenario.code
            scenario_name = scenario.name
            assumptions = scenario.assumptions

        arguments = [
            argument
            for argument in all_arguments
            if argument.id in selected_ids
        ]
        argument_by_id = {argument.id: argument for argument in arguments}
        missing_ids = selected_ids - set(argument_by_id)
        if missing_ids:
            raise ValueError(
                "El escenario contiene argumentos inexistentes o descartados."
            )
        relations = [
            relation
            for relation in all_relations
            if relation.source_argument_id in selected_ids
            and relation.target_argument_id in selected_ids
        ]
        assessments = {
            argument.code: assess_argument_support(argument)
            for argument in arguments
        }
        nodes = [
            ArgumentGraphNode(
                argument_id=argument.id,
                code=argument.code,
                title=argument.title,
                position=argument.position,
                status=argument.status,
                thesis_key=argument.thesis_key,
                thesis_value=argument.thesis_value,
                support_level=assessments[argument.code].level,
                support_score=assessments[argument.code].score,
            )
            for argument in arguments
        ]
        edges = [
            ArgumentGraphEdge(
                relation_id=relation.id,
                code=relation.code,
                source_argument_id=relation.source_argument_id,
                source_code=argument_by_id[relation.source_argument_id].code,
                target_argument_id=relation.target_argument_id,
                target_code=argument_by_id[relation.target_argument_id].code,
                relation_type=relation.relation_type,
                rationale=relation.rationale,
            )
            for relation in relations
        ]
        replied_targets = {
            relation.target_argument_id
            for relation in relations
            if relation.relation_type is ArgumentRelationType.REPLIES
        }
        related_ids = {
            argument_id
            for relation in relations
            for argument_id in (
                relation.source_argument_id,
                relation.target_argument_id,
            )
        }
        unresolved = [
            argument.code
            for argument in arguments
            if argument.position is ArgumentPosition.ADVERSE
            and argument.id not in replied_targets
        ]
        orphan = [
            argument.code
            for argument in arguments
            if argument.id not in related_ids
        ]
        support_total = sum(
            assessment.score for assessment in assessments.values()
        )
        support_average = (
            round(support_total / len(arguments), 2) if arguments else 0.0
        )
        component_count = _connected_component_count(
            selected_ids,
            relations,
        )
        summary: dict[str, int | str | bool | float] = {
            "node_count": len(nodes),
            "edge_count": len(edges),
            "favorable_count": sum(
                node.position is ArgumentPosition.FAVORABLE for node in nodes
            ),
            "adverse_count": sum(
                node.position is ArgumentPosition.ADVERSE for node in nodes
            ),
            "neutral_count": sum(
                node.position is ArgumentPosition.NEUTRAL for node in nodes
            ),
            "support_edge_count": sum(
                edge.relation_type is ArgumentRelationType.SUPPORTS
                for edge in edges
            ),
            "attack_edge_count": sum(
                edge.relation_type is ArgumentRelationType.ATTACKS
                for edge in edges
            ),
            "reply_edge_count": sum(
                edge.relation_type is ArgumentRelationType.REPLIES
                for edge in edges
            ),
            "unresolved_objection_count": len(unresolved),
            "orphan_argument_count": len(orphan),
            "component_count": component_count,
            "support_score_total": support_total,
            "support_score_average": support_average,
            "scenario_status": (
                scenario.status.value if scenario is not None else "Completo"
            ),
            "argumentation_graph_version": "4.1.1",
        }
        warnings: list[str] = []
        if unresolved:
            warnings.append(
                "Objeciones sin réplica: " + ", ".join(unresolved) + "."
            )
        if orphan:
            warnings.append(
                "Argumentos aislados: " + ", ".join(orphan) + "."
            )
        if component_count > 1:
            warnings.append(
                f"El grafo contiene {component_count} componentes desconectados."
            )
        if not any(
            argument.position is ArgumentPosition.FAVORABLE
            for argument in arguments
        ):
            warnings.append("El escenario no contiene argumentos favorables.")
        if not any(
            argument.position is ArgumentPosition.ADVERSE
            for argument in arguments
        ):
            warnings.append("El escenario no contiene argumentos adversos.")
        dot_source = self._build_dot(nodes, edges, scenario_name)
        hash_payload = {
            "case_id": case_id,
            "issue_id": issue_id,
            "scenario": (
                scenario.model_dump(mode="json") if scenario is not None else None
            ),
            "nodes": [node.model_dump(mode="json") for node in nodes],
            "edges": [edge.model_dump(mode="json") for edge in edges],
            "summary": summary,
            "warnings": warnings,
            "assumptions": assumptions,
        }
        input_hash = hashlib.sha256(
            json.dumps(
                hash_payload,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ).encode("utf-8")
        ).hexdigest()
        return ArgumentGraphSnapshot(
            case_id=case_id,
            issue_id=issue_id,
            scenario_id=scenario_id,
            scenario_code=scenario_code,
            scenario_name=scenario_name,
            generated_at=self._repository.snapshot_timestamp(case_id, issue_id),
            input_hash=input_hash,
            nodes=nodes,
            edges=edges,
            summary=summary,
            unresolved_objection_codes=unresolved,
            orphan_argument_codes=orphan,
            warnings=warnings,
            assumptions=assumptions,
            dot_source=dot_source,
        )

    def compare_scenarios(
        self,
        case_id: str,
        issue_id: str,
        first_scenario_id: str | None,
        second_scenario_id: str | None,
    ) -> ScenarioComparison:
        """Compara dos escenarios o cualquiera de ellos con la vista completa."""

        first = self.build_graph(case_id, issue_id, first_scenario_id)
        second = self.build_graph(case_id, issue_id, second_scenario_id)
        first_arguments = {node.code for node in first.nodes}
        second_arguments = {node.code for node in second.nodes}
        first_relations = {edge.code for edge in first.edges}
        second_relations = {edge.code for edge in second.edges}
        metric_keys = (
            "node_count",
            "edge_count",
            "favorable_count",
            "adverse_count",
            "unresolved_objection_count",
            "orphan_argument_count",
            "component_count",
            "support_score_total",
            "support_score_average",
        )
        metric_deltas: dict[str, int | float] = {}
        for key in metric_keys:
            first_value = first.summary[key]
            second_value = second.summary[key]
            if not isinstance(first_value, int | float):
                continue
            if not isinstance(second_value, int | float):
                continue
            metric_deltas[key] = round(second_value - first_value, 2)
        return ScenarioComparison(
            first_scenario_code=first.scenario_code or "COMPLETO",
            second_scenario_code=second.scenario_code or "COMPLETO",
            same_input=first.input_hash == second.input_hash,
            first_input_hash=first.input_hash,
            second_input_hash=second.input_hash,
            added_argument_codes=sorted(second_arguments - first_arguments),
            removed_argument_codes=sorted(first_arguments - second_arguments),
            added_relation_codes=sorted(second_relations - first_relations),
            removed_relation_codes=sorted(first_relations - second_relations),
            metric_deltas=metric_deltas,
        )

    def export_graph_json(
        self,
        case_id: str,
        issue_id: str,
        scenario_id: str | None = None,
    ) -> str:
        """Exporta el grafo o escenario en JSON reproducible."""

        graph = self.build_graph(case_id, issue_id, scenario_id)
        payload = {
            "metadata": self._repository.issue_metadata(case_id, issue_id),
            "graph": graph.model_dump(mode="json"),
            "warning": self._warning_text(),
        }
        return json.dumps(
            payload,
            ensure_ascii=False,
            indent=2,
            sort_keys=True,
        )

    def export_graph_markdown(
        self,
        case_id: str,
        issue_id: str,
        scenario_id: str | None = None,
    ) -> str:
        """Exporta una vista textual del grafo argumental."""

        graph = self.build_graph(case_id, issue_id, scenario_id)
        metadata = self._repository.issue_metadata(case_id, issue_id)
        lines = [
            "# IUS-Razón · Grafo argumental y escenario",
            "",
            f"- Expediente: **{metadata['case_title']}**",
            f"- Problema: `{metadata['issue_code']}` · {metadata['issue_title']}",
            f"- Escenario: **{graph.scenario_code or 'COMPLETO'}** · "
            f"{graph.scenario_name}",
            f"- Huella SHA-256: `{graph.input_hash}`",
            "",
            "## Métricas",
            "",
        ]
        lines.extend(
            f"- {key}: `{value}`"
            for key, value in sorted(graph.summary.items())
        )
        lines.extend(["", "## Supuestos", ""])
        lines.extend(
            f"- {assumption}" for assumption in graph.assumptions
        )
        if not graph.assumptions:
            lines.append("No se registraron supuestos adicionales.")
        lines.extend(["", "## Nodos argumentales", ""])
        for node in graph.nodes:
            lines.append(
                f"- `{node.code}` · {node.title} · {node.position.value} · "
                f"soporte {node.support_level.value} ({node.support_score}/5)."
            )
        if not graph.nodes:
            lines.append("No hay argumentos en esta vista.")
        lines.extend(["", "## Relaciones", ""])
        for edge in graph.edges:
            lines.append(
                f"- `{edge.code}`: {edge.source_code} "
                f"{edge.relation_type.value.lower()} a {edge.target_code}. "
                f"{edge.rationale}"
            )
        if not graph.edges:
            lines.append("No hay relaciones en esta vista.")
        lines.extend(["", "## Advertencias estructurales", ""])
        lines.extend(f"- {warning}" for warning in graph.warnings)
        if not graph.warnings:
            lines.append("No se detectaron advertencias estructurales.")
        lines.extend(["", "## Advertencia jurídica", "", self._warning_text()])
        return "\n".join(lines)

    def export_graph_dot(
        self,
        case_id: str,
        issue_id: str,
        scenario_id: str | None = None,
    ) -> str:
        """Exporta el código DOT para revisión o visualización externa."""

        return self.build_graph(case_id, issue_id, scenario_id).dot_source

    def build_dossier(
        self,
        case_id: str,
        issue_id: str,
    ) -> ArgumentationDossier:
        """Construye una instantánea determinista del debate argumental."""

        self._repository.validate_case_issue(case_id, issue_id)
        arguments = self.list_arguments(case_id, issue_id)
        relations = self.list_relations(case_id, issue_id)
        related_ids: set[str] = set()
        replied_targets: set[str] = set()
        for relation in relations:
            related_ids.add(relation.source_argument_id)
            related_ids.add(relation.target_argument_id)
            if relation.relation_type is ArgumentRelationType.REPLIES:
                replied_targets.add(relation.target_argument_id)

        unresolved = [
            argument.code
            for argument in arguments
            if argument.position is ArgumentPosition.ADVERSE
            and argument.id not in replied_targets
        ]
        orphan = [
            argument.code
            for argument in arguments
            if argument.id not in related_ids
        ]
        support_assessments = [
            assess_argument_support(argument) for argument in arguments
        ]
        missing_information = [
            (
                f"{assessment.argument_code}: faltan "
                f"{', '.join(assessment.missing_components)}."
            )
            for assessment in support_assessments
            if assessment.missing_components
        ]
        missing_information.extend(
            f"{code}: argumento adverso sin réplica registrada."
            for code in unresolved
        )
        favorable_count = sum(
            argument.position is ArgumentPosition.FAVORABLE
            for argument in arguments
        )
        adverse_count = sum(
            argument.position is ArgumentPosition.ADVERSE
            for argument in arguments
        )
        supported_count = sum(
            argument.status
            in {ArgumentStatus.SUPPORTED, ArgumentStatus.VALIDATED}
            for argument in arguments
        )
        attack_count = sum(
            relation.relation_type is ArgumentRelationType.ATTACKS
            for relation in relations
        )
        high_support_count = sum(
            item.level is ArgumentSupportLevel.HIGH
            for item in support_assessments
        )
        summary: dict[str, int | str | bool] = {
            "argument_count": len(arguments),
            "relation_count": len(relations),
            "favorable_count": favorable_count,
            "adverse_count": adverse_count,
            "neutral_count": len(arguments) - favorable_count - adverse_count,
            "supported_count": supported_count,
            "attack_count": attack_count,
            "unresolved_objection_count": len(unresolved),
            "orphan_argument_count": len(orphan),
            "high_support_count": high_support_count,
            "missing_information_count": len(missing_information),
            "has_unresolved_objections": bool(unresolved),
            "scenario_count": len(self.list_scenarios(case_id, issue_id)),
            "argumentation_version": "4.1.1",
        }
        payload = {
            "case_id": case_id,
            "issue_id": issue_id,
            "arguments": [
                argument.model_dump(mode="json") for argument in arguments
            ],
            "relations": [
                relation.model_dump(mode="json") for relation in relations
            ],
            "support_assessments": [
                item.model_dump(mode="json") for item in support_assessments
            ],
            "summary": summary,
            "unresolved_objection_codes": unresolved,
            "orphan_argument_codes": orphan,
            "missing_information": missing_information,
        }
        input_hash = hashlib.sha256(
            json.dumps(
                payload,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ).encode("utf-8")
        ).hexdigest()
        return ArgumentationDossier(
            case_id=case_id,
            issue_id=issue_id,
            generated_at=self._repository.snapshot_timestamp(case_id, issue_id),
            input_hash=input_hash,
            arguments=arguments,
            relations=relations,
            support_assessments=support_assessments,
            summary=summary,
            unresolved_objection_codes=unresolved,
            orphan_argument_codes=orphan,
            missing_information=missing_information,
        )


    def export_json(self, case_id: str, issue_id: str) -> str:
        """Exporta un expediente argumental reproducible en JSON."""

        dossier = self.build_dossier(case_id, issue_id)
        metadata = self._repository.issue_metadata(case_id, issue_id)
        payload = {
            "metadata": metadata,
            "dossier": dossier.model_dump(mode="json"),
            "warning": self._warning_text(),
        }
        return json.dumps(
            payload,
            ensure_ascii=False,
            indent=2,
            sort_keys=True,
        )

    def export_markdown(self, case_id: str, issue_id: str) -> str:
        """Genera un informe argumental en Markdown."""

        dossier = self.build_dossier(case_id, issue_id)
        metadata = self._repository.issue_metadata(case_id, issue_id)
        relation_labels = self._relation_labels(dossier.arguments)
        support_by_code = {
            item.argument_code: item for item in dossier.support_assessments
        }
        lines = [
            "# IUS-Razón · Informe de argumentación jurídica",
            "",
            f"- Expediente: **{metadata['case_title']}**",
            f"- Problema: `{metadata['issue_code']}` · {metadata['issue_title']}",
            f"- Pregunta: {metadata['question']}",
            f"- Jurisdicción registrada: {metadata['jurisdiction']}",
            f"- Huella SHA-256: `{dossier.input_hash}`",
            "",
            "## Resumen",
            "",
        ]
        for key, value in sorted(dossier.summary.items()):
            lines.append(f"- {key}: `{value}`")
        lines.extend(["", "## Argumentos", ""])
        if not dossier.arguments:
            lines.append("No hay argumentos registrados.")
        for argument in dossier.arguments:
            assessment = support_by_code[argument.code]
            missing = (
                ", ".join(assessment.missing_components)
                if assessment.missing_components
                else "Ninguno"
            )
            lines.extend(
                [
                    f"### {argument.code} · {argument.title}",
                    "",
                    f"- Posición: **{argument.position.value}**",
                    f"- Estado: {argument.status.value}",
                    (
                        f"- Tesis: `{argument.thesis_key}` = "
                        f"**{argument.thesis_value.value}**"
                    ),
                    (
                        "- Soporte descriptivo: "
                        f"**{assessment.level.value}** ({assessment.score}/5)"
                    ),
                    f"- Información faltante: {missing}",
                    f"- Conclusión vinculada: {argument.conclusion_id or 'Ninguna'}",
                    f"- Hechos: {', '.join(argument.fact_codes) or 'Ninguno'}",
                    (
                        "- Pruebas: "
                        f"{', '.join(argument.evidence_codes) or 'Ninguna'}"
                    ),
                    f"- Fuentes: {', '.join(argument.source_codes) or 'Ninguna'}",
                    f"- Reglas: {', '.join(argument.rule_codes) or 'Ninguna'}",
                    (
                        "- Premisas: "
                        f"{', '.join(argument.assertion_codes) or 'Ninguna'}"
                    ),
                    "",
                    "#### Pretensión",
                    "",
                    argument.claim,
                    "",
                    "#### Razonamiento",
                    "",
                    argument.reasoning,
                    "",
                ]
            )
        lines.extend(["## Relaciones argumentales", ""])
        if not dossier.relations:
            lines.append("No hay relaciones registradas.")
        for relation in dossier.relations:
            source = relation_labels.get(
                relation.source_argument_id,
                relation.source_argument_id,
            )
            target = relation_labels.get(
                relation.target_argument_id,
                relation.target_argument_id,
            )
            lines.extend(
                [
                    (
                        f"- `{relation.code}`: **{source}** "
                        f"{relation.relation_type.value.lower()} a "
                        f"**{target}**."
                    ),
                    f"  - Razón: {relation.rationale}",
                ]
            )
        lines.extend(
            [
                "",
                "## Información faltante",
                "",
            ]
        )
        if dossier.missing_information:
            lines.extend(
                f"- {item}" for item in dossier.missing_information
            )
        else:
            lines.append(
                "No se detectaron vacíos estructurales en la trazabilidad."
            )
        lines.extend(
            [
                "",
                "## Objeciones pendientes",
                "",
                (
                    ", ".join(dossier.unresolved_objection_codes)
                    if dossier.unresolved_objection_codes
                    else "No se detectaron argumentos adversos sin réplica."
                ),
                "",
                "## Argumentos aislados",
                "",
                (
                    ", ".join(dossier.orphan_argument_codes)
                    if dossier.orphan_argument_codes
                    else "No se detectaron argumentos sin relaciones."
                ),
                "",
                "## Advertencia",
                "",
                self._warning_text(),
                "",
            ]
        )
        return "\n".join(lines)


    def export_docx(self, case_id: str, issue_id: str) -> bytes:
        """Genera un DOCX local con la misma instantánea argumental."""

        dossier = self.build_dossier(case_id, issue_id)
        metadata = self._repository.issue_metadata(case_id, issue_id)
        relation_labels = self._relation_labels(dossier.arguments)
        support_by_code = {
            item.argument_code: item for item in dossier.support_assessments
        }
        document = Document()
        document.add_heading("IUS-Razón · Informe de argumentación jurídica", 0)
        document.add_paragraph(
            f"Expediente: {metadata['case_title']}\n"
            f"Problema: {metadata['issue_code']} · {metadata['issue_title']}\n"
            f"Pregunta: {metadata['question']}\n"
            f"Huella SHA-256: {dossier.input_hash}"
        )
        document.add_heading("Resumen", level=1)
        for key, value in sorted(dossier.summary.items()):
            document.add_paragraph(f"{key}: {value}", style="List Bullet")

        document.add_heading("Argumentos", level=1)
        if not dossier.arguments:
            document.add_paragraph("No hay argumentos registrados.")
        for argument in dossier.arguments:
            assessment = support_by_code[argument.code]
            missing = (
                ", ".join(assessment.missing_components)
                if assessment.missing_components
                else "Ninguno"
            )
            document.add_heading(
                f"{argument.code} · {argument.title}",
                level=2,
            )
            document.add_paragraph(
                f"Posición: {argument.position.value}\n"
                f"Estado: {argument.status.value}\n"
                f"Tesis: {argument.thesis_key} = {argument.thesis_value.value}\n"
                f"Soporte descriptivo: {assessment.level.value} "
                f"({assessment.score}/5)\n"
                f"Información faltante: {missing}"
            )
            document.add_heading("Pretensión", level=3)
            document.add_paragraph(argument.claim)
            document.add_heading("Razonamiento", level=3)
            document.add_paragraph(argument.reasoning)
            document.add_paragraph(
                "Hechos: "
                f"{', '.join(argument.fact_codes) or 'Ninguno'}\n"
                "Pruebas: "
                f"{', '.join(argument.evidence_codes) or 'Ninguna'}\n"
                "Fuentes: "
                f"{', '.join(argument.source_codes) or 'Ninguna'}\n"
                "Reglas: "
                f"{', '.join(argument.rule_codes) or 'Ninguna'}\n"
                "Premisas: "
                f"{', '.join(argument.assertion_codes) or 'Ninguna'}"
            )

        document.add_heading("Relaciones argumentales", level=1)
        if not dossier.relations:
            document.add_paragraph("No hay relaciones registradas.")
        for relation in dossier.relations:
            source = relation_labels.get(
                relation.source_argument_id,
                relation.source_argument_id,
            )
            target = relation_labels.get(
                relation.target_argument_id,
                relation.target_argument_id,
            )
            document.add_paragraph(
                f"{relation.code}: {source} "
                f"{relation.relation_type.value.lower()} a {target}. "
                f"{relation.rationale}",
                style="List Bullet",
            )

        document.add_heading("Información faltante", level=1)
        if dossier.missing_information:
            for item in dossier.missing_information:
                document.add_paragraph(item, style="List Bullet")
        else:
            document.add_paragraph(
                "No se detectaron vacíos estructurales en la trazabilidad."
            )
        document.add_heading("Objeciones pendientes", level=1)
        document.add_paragraph(
            ", ".join(dossier.unresolved_objection_codes)
            if dossier.unresolved_objection_codes
            else "No se detectaron argumentos adversos sin réplica."
        )
        document.add_heading("Advertencia", level=1)
        document.add_paragraph(self._warning_text())
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()


    def _validate_scenario_arguments(
        self,
        case_id: str,
        issue_id: str,
        argument_ids: list[str],
    ) -> None:
        """Impide referencias cruzadas o inexistentes en escenarios."""

        known = {
            argument.id
            for argument in self.list_arguments(
                case_id,
                issue_id,
                include_discarded=True,
            )
        }
        missing = [argument_id for argument_id in argument_ids if argument_id not in known]
        if missing:
            raise ValueError(
                "Los argumentos del escenario no pertenecen al problema "
                "jurídico activo."
            )

    @staticmethod
    def _build_dot(
        nodes: list[ArgumentGraphNode],
        edges: list[ArgumentGraphEdge],
        title: str,
    ) -> str:
        """Construye DOT determinista sin HTML ni recursos externos."""

        lines = [
            "digraph argumentacion {",
            '  graph [rankdir="LR", labelloc="t", label="'
            + _escape_dot(title)
            + '"];',
            '  node [fontname="Arial"];',
            '  edge [fontname="Arial"];',
        ]
        shape_by_position = {
            ArgumentPosition.FAVORABLE: "box",
            ArgumentPosition.ADVERSE: "octagon",
            ArgumentPosition.NEUTRAL: "ellipse",
        }
        for node in sorted(nodes, key=lambda item: item.code):
            label = (
                f"{node.code}\\n{node.title}\\n"
                f"{node.position.value} · {node.support_level.value}"
            )
            lines.append(
                f'  "{_escape_dot(node.code)}" '
                f'[shape="{shape_by_position[node.position]}", '
                f'label="{_escape_dot(label)}"];'
            )
        style_by_relation = {
            ArgumentRelationType.SUPPORTS: ("solid", "normal"),
            ArgumentRelationType.ATTACKS: ("bold", "tee"),
            ArgumentRelationType.REPLIES: ("dashed", "normal"),
        }
        for edge in sorted(edges, key=lambda item: item.code):
            style, arrowhead = style_by_relation[edge.relation_type]
            lines.append(
                f'  "{_escape_dot(edge.source_code)}" -> '
                f'"{_escape_dot(edge.target_code)}" '
                f'[label="{_escape_dot(edge.relation_type.value)}", '
                f'style="{style}", arrowhead="{arrowhead}"];'
            )
        lines.append("}")
        return "\n".join(lines)

    def _validate_argument_codes(
        self,
        case_id: str,
        issue_id: str,
        payload: LegalArgumentCreate | LegalArgumentUpdate,
    ) -> None:
        groups = (
            ("hechos", payload.fact_codes, ("H-",)),
            ("pruebas", payload.evidence_codes, ("P-",)),
            ("fuentes", payload.source_codes, ("N-", "J-", "D-")),
            ("reglas", payload.rule_codes, ("R-",)),
            ("premisas", payload.assertion_codes, ("A-",)),
        )
        known = self._repository.known_codes(case_id, issue_id)
        for label, codes, prefixes in groups:
            invalid_prefix = [
                code for code in codes if not code.startswith(prefixes)
            ]
            if invalid_prefix:
                raise ValueError(
                    f"Los códigos de {label} tienen un prefijo inválido: "
                    f"{', '.join(invalid_prefix)}."
                )
            missing = [code for code in codes if code not in known]
            if missing:
                raise ValueError(
                    f"Los códigos de {label} no existen en el expediente: "
                    f"{', '.join(missing)}."
                )

    def _validate_conclusion_link(
        self,
        payload: LegalArgumentCreate,
    ) -> None:
        if payload.conclusion_id is None:
            return
        context = self._repository.conclusion_context(payload.conclusion_id)
        if (
            str(context["case_id"]) != payload.case_id
            or str(context["issue_id"]) != payload.issue_id
        ):
            raise ValueError(
                "La conclusión vinculada pertenece a otro problema jurídico."
            )

    def _backup_before_mutation(self, operation: str) -> None:
        self._mutation_backup.before_mutation()
        LOGGER.info(
            "Política de respaldo previo completada. operation=%s",
            operation,
        )

    @staticmethod
    def _relation_labels(
        arguments: list[LegalArgumentRecord],
    ) -> dict[str, str]:
        return {
            argument.id: f"{argument.code} · {argument.title}"
            for argument in arguments
        }

    @staticmethod
    def _warning_text() -> str:
        return (
            "Este informe organiza argumentos registrados por el usuario. "
            "No verifica automáticamente autenticidad, vigencia, aplicabilidad "
            "ni suficiencia jurídica y no constituye asesoría legal."
        )
