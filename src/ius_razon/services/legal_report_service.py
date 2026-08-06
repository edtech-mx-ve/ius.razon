from __future__ import annotations

import hashlib
import io
import json
from collections.abc import Iterable
from datetime import datetime

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import _Row
from docx.text.paragraph import Paragraph

from ius_razon.domain.argumentation_models import (
    ArgumentationDossier,
    ArgumentGraphSnapshot,
    LegalArgumentRecord,
    ScenarioComparison,
)
from ius_razon.domain.reasoning_models import ReasoningRunReport
from ius_razon.domain.report_models import (
    IntegralLegalReport,
    IntegralReportRequest,
    ScenarioNarrative,
    TraceabilityRow,
)
from ius_razon.services.argumentation_service import ArgumentationService
from ius_razon.services.case_service import CaseService
from ius_razon.services.reasoning_service import ReasoningService

REPORT_VERSION = "4.2.1"


def _clean_lines(values: Iterable[str]) -> list[str]:
    """Normaliza una secuencia preservando orden y eliminando duplicados."""

    normalized: list[str] = []
    for value in values:
        item = value.strip()
        if item and item not in normalized:
            normalized.append(item)
    return normalized


def _codes(values: Iterable[str]) -> str:
    """Convierte códigos en texto legible para reportes."""

    items = list(values)
    return ", ".join(items) if items else "Ninguno"


def _count_phrase(
    count: int,
    singular: str,
    plural: str | None = None,
) -> str:
    """Construye una frase cuantificada con concordancia básica."""

    plural_form = plural or f"{singular}s"
    noun = singular if count == 1 else plural_form
    return f"{count} {noun}"


def _ensure_sentence(value: str) -> str:
    """Normaliza el cierre de una oración sin duplicar puntuación."""

    normalized = value.strip()
    if not normalized:
        return normalized
    if normalized.endswith((".", "!", "?")):
        return normalized
    return normalized + "."


def _argument_origin(argument: LegalArgumentRecord) -> str:
    """Describe el origen trazable de un argumento."""

    if argument.conclusion_id is None:
        return "Manual"
    return "Derivado de una conclusión del motor"


def _latest_datetime(values: Iterable[datetime | None]) -> datetime:
    """Obtiene la marca temporal más reciente de una instantánea."""

    available = [value for value in values if value is not None]
    if not available:
        raise ValueError("No fue posible determinar la fecha de la instantánea.")
    return max(available)


def _mark_table_header(row: _Row) -> None:
    """Marca una fila como encabezado repetible y accesible."""

    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def _add_page_number(paragraph: Paragraph) -> None:
    """Inserta un campo PAGE en un párrafo de python-docx."""

    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


class LegalReportService:
    """Construye un informe jurídico integral a partir de datos trazables."""

    def __init__(
        self,
        case_service: CaseService,
        reasoning_service: ReasoningService,
        argumentation_service: ArgumentationService,
    ) -> None:
        self._case_service = case_service
        self._reasoning_service = reasoning_service
        self._argumentation_service = argumentation_service

    def build_report(self, request: IntegralReportRequest) -> IntegralLegalReport:
        """Construye una instantánea integral, determinista y verificable."""

        case_summary = self._case_service.get_case_summary(request.case_id)
        issues = self._case_service.list_legal_issues(request.case_id)
        issue = next(
            (item for item in issues if item.id == request.issue_id),
            None,
        )
        if issue is None:
            raise ValueError("El problema jurídico no pertenece al expediente activo.")

        reasoning = self._reasoning_service.get_run_report(
            request.reasoning_run_id
        )
        if (
            reasoning.run.case_id != request.case_id
            or reasoning.run.issue_id != request.issue_id
        ):
            raise ValueError(
                "La ejecución seleccionada pertenece a otro problema jurídico."
            )

        parties = self._case_service.list_parties(request.case_id)
        facts = self._case_service.list_facts(request.case_id)
        evidence = self._case_service.list_evidence(request.case_id)
        fact_evidence_links = self._case_service.list_fact_evidence_links(
            request.case_id
        )
        issue_source_links = [
            link
            for link in self._case_service.list_issue_source_links(
                request.case_id
            )
            if link.issue_id == request.issue_id
        ]
        linked_codes = {link.source_code for link in issue_source_links}
        norms = [
            item
            for item in self._case_service.list_norms(request.case_id)
            if item.code in linked_codes
        ]
        jurisprudence = [
            item
            for item in self._case_service.list_jurisprudence(request.case_id)
            if item.code in linked_codes
        ]
        doctrine = [
            item
            for item in self._case_service.list_doctrine(request.case_id)
            if item.code in linked_codes
        ]

        argumentation = self._argumentation_service.build_dossier(
            request.case_id,
            request.issue_id,
        )
        base_graph = self._argumentation_service.build_graph(
            request.case_id,
            request.issue_id,
            request.base_scenario_id,
        )
        compared_graph: ArgumentGraphSnapshot | None = None
        comparison: ScenarioComparison | None = None
        narrative: ScenarioNarrative | None = None
        if request.compared_scenario_id is not None:
            compared_graph = self._argumentation_service.build_graph(
                request.case_id,
                request.issue_id,
                request.compared_scenario_id,
            )
            comparison = self._argumentation_service.compare_scenarios(
                request.case_id,
                request.issue_id,
                request.base_scenario_id,
                request.compared_scenario_id,
            )
            narrative = self._scenario_narrative(
                base_graph,
                compared_graph,
                comparison,
            )

        traceability_rows = self._traceability_rows(argumentation, reasoning)
        executive_summary = (
            request.executive_summary
            or self._automatic_executive_summary(
                issue_code=issue.code,
                reasoning=reasoning,
                argument_count=len(argumentation.arguments),
                unresolved_count=len(
                    argumentation.unresolved_objection_codes
                ),
                comparison=narrative,
            )
        )
        methodology = [
            (
                "El informe integra datos registrados por el usuario en el "
                "expediente, sin consultar fuentes externas."
            ),
            (
                "Las conclusiones de inferencia proceden del motor determinista "
                f"{reasoning.run.engine_version} y conservan su traza."
            ),
            (
                "La evaluación de soporte argumental mide completitud de "
                "trazabilidad, no probabilidad de éxito."
            ),
            (
                "Los escenarios comparan subconjuntos de argumentos y "
                "relaciones bajo supuestos explícitos."
            ),
        ]
        findings = self._findings(
            reasoning=reasoning,
            argumentation=argumentation,
            base_graph=base_graph,
            narrative=narrative,
        )
        limitations = self._limitations(
            argumentation.missing_information,
            argumentation.unresolved_objection_codes,
            request.additional_limitations,
        )
        recommendations = self._recommendations(
            request.recommendations,
            limitations,
        )
        section_index = [
            "1. Identificación y alcance",
            "2. Resumen ejecutivo",
            "3. Metodología y advertencias",
            "4. Antecedentes del expediente",
            "5. Hechos y pruebas",
            "6. Fuentes jurídicas vinculadas",
            "7. Inferencia y trazabilidad",
            "8. Argumentación jurídica",
            "9. Escenarios alternativos",
            "10. Hallazgos y conclusiones",
            "11. Limitaciones e información faltante",
            "12. Recomendaciones de revisión",
            "Anexo A. Matriz de trazabilidad",
            "Anexo B. Huellas y reproducibilidad",
        ]
        generated_at = _latest_datetime(
            [
                case_summary.case.updated_at,
                issue.updated_at,
                reasoning.run.completed_at,
                reasoning.run.started_at,
                argumentation.generated_at,
                base_graph.generated_at,
                compared_graph.generated_at if compared_graph else None,
            ]
        )

        payload = {
            "report_version": REPORT_VERSION,
            "request": request.model_dump(mode="json"),
            "case": case_summary.case.model_dump(mode="json"),
            "issue": issue.model_dump(mode="json"),
            "parties": [item.model_dump(mode="json") for item in parties],
            "facts": [item.model_dump(mode="json") for item in facts],
            "evidence": [item.model_dump(mode="json") for item in evidence],
            "fact_evidence_links": [
                item.model_dump(mode="json") for item in fact_evidence_links
            ],
            "norms": [item.model_dump(mode="json") for item in norms],
            "jurisprudence": [
                item.model_dump(mode="json") for item in jurisprudence
            ],
            "doctrine": [item.model_dump(mode="json") for item in doctrine],
            "issue_source_links": [
                item.model_dump(mode="json") for item in issue_source_links
            ],
            "reasoning": reasoning.model_dump(mode="json"),
            "argumentation": argumentation.model_dump(mode="json"),
            "base_graph": base_graph.model_dump(mode="json"),
            "compared_graph": (
                compared_graph.model_dump(mode="json")
                if compared_graph is not None
                else None
            ),
            "scenario_comparison": (
                comparison.model_dump(mode="json")
                if comparison is not None
                else None
            ),
            "executive_summary": executive_summary,
            "methodology": methodology,
            "findings": findings,
            "limitations": limitations,
            "recommendations": recommendations,
            "traceability_rows": [
                item.model_dump(mode="json") for item in traceability_rows
            ],
            "section_index": section_index,
        }
        input_hash = hashlib.sha256(
            json.dumps(
                payload,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ).encode("utf-8")
        ).hexdigest()

        return IntegralLegalReport(
            report_version=REPORT_VERSION,
            generated_at=generated_at,
            input_hash=input_hash,
            request=request,
            case=case_summary.case,
            issue=issue,
            parties=parties,
            facts=facts,
            evidence=evidence,
            fact_evidence_links=fact_evidence_links,
            norms=norms,
            jurisprudence=jurisprudence,
            doctrine=doctrine,
            issue_source_links=issue_source_links,
            reasoning=reasoning,
            argumentation=argumentation,
            base_graph=base_graph,
            compared_graph=compared_graph,
            scenario_comparison=comparison,
            scenario_narrative=narrative,
            executive_summary=executive_summary,
            methodology=methodology,
            findings=findings,
            limitations=limitations,
            recommendations=recommendations,
            traceability_rows=traceability_rows,
            section_index=section_index,
            warning=self._warning_text(),
        )

    def export_json(self, request: IntegralReportRequest) -> str:
        """Exporta el informe integral en JSON legible."""

        report = self.build_report(request)
        return json.dumps(
            report.model_dump(mode="json"),
            ensure_ascii=False,
            indent=2,
            sort_keys=True,
        )

    def export_markdown(self, request: IntegralReportRequest) -> str:
        """Exporta el informe integral en Markdown."""

        report = self.build_report(request)
        lines = [
            f"# {report.request.title}",
            "",
            f"- Expediente: **{report.case.title}**",
            f"- Problema: `{report.issue.code}` · {report.issue.title}",
            f"- Jurisdicción: {report.case.jurisdiction}",
            f"- Fecha de instantánea: `{report.generated_at.isoformat()}`",
            f"- Versión del informe: `{report.report_version}`",
            f"- Huella SHA-256: `{report.input_hash}`",
            "",
            "## Índice",
            "",
        ]
        lines.extend(f"- {item}" for item in report.section_index)
        lines.extend(
            [
                "",
                "## 1. Identificación y alcance",
                "",
                f"**Objeto:** {report.request.purpose}",
                "",
                f"**Pregunta jurídica:** {report.issue.question}",
                "",
                "### Partes registradas",
                "",
            ]
        )
        if report.parties:
            for party in report.parties:
                lines.append(
                    f"- {party.name_alias} · {party.legal_role.value} · "
                    f"{party.party_type.value}."
                )
        else:
            lines.append("No hay partes registradas.")
        lines.extend(
            [
                "",
                "## 2. Resumen ejecutivo",
                "",
                report.executive_summary,
                "",
                "## 3. Metodología y advertencias",
                "",
            ]
        )
        lines.extend(f"- {item}" for item in report.methodology)
        lines.extend(["", report.warning, "", "## 4. Antecedentes del expediente", ""])
        lines.extend(
            [
                f"- Materia: {report.case.matter}",
                f"- Estado: {report.case.status.value}",
                f"- Confidencialidad: {report.case.confidentiality.value}",
                f"- Objetivo registrado: {report.case.objective or 'No registrado'}",
                f"- Descripción: {report.case.description}",
                "",
                "## 5. Hechos y pruebas",
                "",
                "### Hechos",
                "",
            ]
        )
        if report.facts:
            for fact in report.facts:
                lines.append(
                    f"- `{fact.code}` · {fact.status.value} · {fact.description}"
                )
        else:
            lines.append("No hay hechos registrados.")
        lines.extend(["", "### Pruebas", ""])
        if report.evidence:
            for item in report.evidence:
                lines.append(
                    f"- `{item.code}` · {item.evidence_type.value} · "
                    f"{item.evaluation_status.value} · {item.description}"
                )
        else:
            lines.append("No hay pruebas registradas.")
        lines.extend(["", "## 6. Fuentes jurídicas vinculadas", ""])
        self._append_markdown_sources(lines, report)
        lines.extend(["", "## 7. Inferencia y trazabilidad", ""])
        lines.extend(
            [
                f"- Ejecución: `{report.reasoning.run.id}`",
                f"- Motor: `{report.reasoning.run.engine_version}`",
                f"- Huella de entrada: `{report.reasoning.run.input_hash}`",
                f"- Estado: {report.reasoning.run.status.value}",
                "",
                "### Conclusiones",
                "",
            ]
        )
        if report.reasoning.conclusions:
            for conclusion in report.reasoning.conclusions:
                lines.extend(
                    [
                        f"#### {conclusion.code} · {conclusion.predicate_key}",
                        "",
                        conclusion.statement,
                        "",
                        f"- Valor: **{conclusion.value.value}**",
                        f"- Estado: {conclusion.status.value}",
                        f"- Soporte: {conclusion.support_level.value}",
                        f"- Reglas: {_codes(conclusion.rule_codes)}",
                        (
                            "- Premisas: "
                            f"{_codes(conclusion.supporting_assertion_codes)}"
                        ),
                        f"- Fuentes: {_codes(conclusion.source_codes)}",
                        "",
                    ]
                )
        else:
            lines.append("No se derivaron conclusiones en la ejecución seleccionada.")
        lines.extend(["", "## 8. Argumentación jurídica", ""])
        self._append_markdown_arguments(lines, report)
        lines.extend(["", "## 9. Escenarios alternativos", ""])
        self._append_markdown_scenarios(lines, report)
        lines.extend(["", "## 10. Hallazgos y conclusiones", ""])
        lines.extend(f"- {item}" for item in report.findings)
        if report.request.analyst_conclusions:
            lines.extend(["", "### Conclusiones registradas por el analista", ""])
            lines.extend(
                f"- {item}" for item in report.request.analyst_conclusions
            )
        lines.extend(
            [
                "",
                "## 11. Limitaciones e información faltante",
                "",
            ]
        )
        lines.extend(f"- {item}" for item in report.limitations)
        lines.extend(["", "## 12. Recomendaciones de revisión", ""])
        lines.extend(f"- {item}" for item in report.recommendations)
        if report.request.include_traceability_appendix:
            lines.extend(["", "## Anexo A. Matriz de trazabilidad", ""])
            for row in report.traceability_rows:
                lines.extend(
                    [
                        f"### {row.argument_code} · {row.position}",
                        "",
                        f"- Tesis: `{row.thesis}`",
                        f"- Origen: {row.origin}",
                        (
                            "- Conclusión: "
                            f"{row.conclusion_code or 'No aplica (argumento manual)'}"
                        ),
                        (
                            f"- Soporte: {row.support_level} "
                            f"({row.support_score}/5)"
                        ),
                        f"- Hechos: {_codes(row.fact_codes)}",
                        f"- Pruebas: {_codes(row.evidence_codes)}",
                        f"- Fuentes: {_codes(row.source_codes)}",
                        f"- Reglas: {_codes(row.rule_codes)}",
                        f"- Premisas: {_codes(row.assertion_codes)}",
                        "",
                    ]
                )
        lines.extend(
            [
                "## Anexo B. Huellas y reproducibilidad",
                "",
                f"- Informe: `{report.input_hash}`",
                f"- Inferencia: `{report.reasoning.run.input_hash}`",
                f"- Grafo base: `{report.base_graph.input_hash}`",
            ]
        )
        if report.compared_graph is not None:
            lines.append(
                f"- Grafo comparado: `{report.compared_graph.input_hash}`"
            )
        lines.extend(["", "## Advertencia final", "", report.warning, ""])
        return "\n".join(lines)

    def export_docx(self, request: IntegralReportRequest) -> bytes:
        """Genera un DOCX integral con estructura, tablas y trazabilidad."""

        report = self.build_report(request)
        document = Document()
        self._configure_document(document, report)
        self._add_cover(document, report)
        document.add_page_break()  # type: ignore[no-untyped-call]
        self._add_contents(document, report)
        document.add_page_break()  # type: ignore[no-untyped-call]
        self._add_identification(document, report)
        self._add_executive_summary(document, report)
        self._add_methodology(document, report)
        self._add_background(document, report)
        self._add_facts_and_evidence(document, report)
        self._add_sources(document, report)
        self._add_reasoning(document, report)
        self._add_argumentation(document, report)
        self._add_scenarios(document, report)
        self._add_findings(document, report)
        self._add_limitations(document, report)
        self._add_recommendations(document, report)
        if report.request.include_traceability_appendix:
            self._add_traceability_appendix(document, report)
        self._add_reproducibility_appendix(document, report)
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()

    @staticmethod
    def _automatic_executive_summary(
        *,
        issue_code: str,
        reasoning: ReasoningRunReport,
        argument_count: int,
        unresolved_count: int,
        comparison: ScenarioNarrative | None,
    ) -> str:
        conclusions = reasoning.conclusions
        if conclusions:
            conclusion_text = "; ".join(
                (
                    f"{item.predicate_key}={item.value.value} "
                    f"({item.status.value}, soporte {item.support_level.value})"
                )
                for item in conclusions
            )
        else:
            conclusion_text = "la ejecución seleccionada no derivó conclusiones"
        argument_phrase = _count_phrase(argument_count, "argumento")
        objection_phrase = _count_phrase(
            unresolved_count,
            "objeción sin réplica",
            "objeciones sin réplica",
        )
        summary = (
            f"El informe integra el análisis del problema {issue_code}. "
            f"El motor determinista registró {conclusion_text}. "
            f"La red argumental contiene {argument_phrase} y "
            f"{objection_phrase}."
        )
        if comparison is not None:
            summary += " " + " ".join(comparison.statements)
        return summary

    @staticmethod
    def _scenario_narrative(
        base: ArgumentGraphSnapshot,
        compared: ArgumentGraphSnapshot,
        comparison: ScenarioComparison,
    ) -> ScenarioNarrative:
        statements: list[str] = []
        if comparison.added_argument_codes:
            statements.append(
                "El escenario comparado incorpora los argumentos "
                + ", ".join(comparison.added_argument_codes)
                + "."
            )
        if comparison.removed_argument_codes:
            statements.append(
                "El escenario comparado excluye los argumentos "
                + ", ".join(comparison.removed_argument_codes)
                + "."
            )
        if comparison.added_relation_codes:
            statements.append(
                "Se incorporan las relaciones "
                + ", ".join(comparison.added_relation_codes)
                + "."
            )
        if comparison.removed_relation_codes:
            statements.append(
                "Se excluyen las relaciones "
                + ", ".join(comparison.removed_relation_codes)
                + "."
            )
        unresolved_delta = comparison.metric_deltas.get(
            "unresolved_objection_count",
            0,
        )
        if unresolved_delta < 0:
            statements.append(
                "Las objeciones pendientes disminuyen en "
                f"{abs(int(unresolved_delta))}."
            )
        elif unresolved_delta > 0:
            statements.append(
                "Las objeciones pendientes aumentan en "
                f"{int(unresolved_delta)}."
            )
        if not statements:
            statements.append(
                "No se detectaron diferencias estructurales entre los escenarios."
            )
        return ScenarioNarrative(
            base_label=f"{base.scenario_code or 'COMPLETO'} · {base.scenario_name}",
            compared_label=(
                f"{compared.scenario_code or 'COMPLETO'} · "
                f"{compared.scenario_name}"
            ),
            statements=statements,
        )

    @staticmethod
    def _findings(
        *,
        reasoning: ReasoningRunReport,
        argumentation: ArgumentationDossier,
        base_graph: ArgumentGraphSnapshot,
        narrative: ScenarioNarrative | None,
    ) -> list[str]:
        findings: list[str] = []
        if reasoning.conclusions:
            for conclusion in reasoning.conclusions:
                findings.append(
                    f"{conclusion.code} establece "
                    f"{conclusion.predicate_key}={conclusion.value.value} "
                    f"con carácter {conclusion.status.value.lower()} y soporte "
                    f"{conclusion.support_level.value.lower()}."
                )
        else:
            findings.append(
                "La ejecución seleccionada no produjo conclusiones vigentes."
            )
        favorable_count = int(base_graph.summary["favorable_count"])
        adverse_count = int(base_graph.summary["adverse_count"])
        unresolved_count = int(
            base_graph.summary["unresolved_objection_count"]
        )
        favorable_phrase = _count_phrase(
            favorable_count,
            "argumento favorable",
            "argumentos favorables",
        )
        adverse_phrase = _count_phrase(
            adverse_count,
            "argumento adverso",
            "argumentos adversos",
        )
        unresolved_phrase = _count_phrase(
            unresolved_count,
            "objeción pendiente",
            "objeciones pendientes",
        )
        findings.append(
            "La red argumental base contiene "
            f"{favorable_phrase} y {adverse_phrase}; registra "
            f"{unresolved_phrase}."
        )
        manual_arguments = [
            argument.code
            for argument in argumentation.arguments
            if argument.conclusion_id is None
        ]
        if manual_arguments:
            if len(manual_arguments) == 1:
                findings.append(
                    f"{manual_arguments[0]} es un argumento manual sin "
                    "conclusión inferida asociada; su tesis y trazabilidad "
                    "fueron registradas por el usuario."
                )
            else:
                findings.append(
                    f"{', '.join(manual_arguments)} son argumentos manuales "
                    "sin conclusiones inferidas asociadas; sus tesis y "
                    "trazabilidad fueron registradas por el usuario."
                )
        if argumentation.missing_information:
            findings.append(
                "Persisten vacíos estructurales de trazabilidad que requieren "
                "revisión humana."
            )
        if narrative is not None:
            findings.extend(narrative.statements)
        return _clean_lines(findings)

    @staticmethod
    def _limitations(
        missing_information: list[str],
        unresolved_codes: list[str],
        additional: list[str],
    ) -> list[str]:
        limitations = [
            (
                "El sistema no verifica automáticamente autenticidad, vigencia, "
                "jerarquía ni aplicabilidad jurídica de las fuentes."
            ),
            (
                "Las prioridades, premisas, reglas, argumentos y escenarios "
                "dependen de datos registrados y validados por el usuario."
            ),
            (
                "El nivel de soporte es descriptivo y no equivale a probabilidad "
                "de éxito, sentencia ni asesoría legal."
            ),
        ]
        limitations.extend(missing_information)
        if unresolved_codes:
            limitations.append(
                "Existen objeciones sin réplica: "
                + ", ".join(unresolved_codes)
                + "."
            )
        limitations.extend(additional)
        return _clean_lines(limitations)

    @staticmethod
    def _recommendations(
        requested: list[str],
        limitations: list[str],
    ) -> list[str]:
        if requested:
            return _clean_lines(requested)
        recommendations = [
            (
                "Verificar documentalmente la autenticidad, vigencia y "
                "aplicabilidad de cada fuente antes de utilizar el informe."
            ),
            (
                "Revisar las premisas con valor desconocido y documentar toda "
                "modificación antes de ejecutar nuevamente el motor."
            ),
            (
                "Resolver los vacíos de trazabilidad y conservar la revisión "
                "humana del informe junto con su huella SHA-256."
            ),
        ]
        if any("objeciones sin réplica" in item.lower() for item in limitations):
            recommendations.append(
                "Registrar una réplica sustentada o mantener expresamente la "
                "objeción como cuestión no resuelta."
            )
        return recommendations

    @staticmethod
    def _traceability_rows(
        argumentation: ArgumentationDossier,
        reasoning: ReasoningRunReport,
    ) -> list[TraceabilityRow]:
        support_by_code = {
            item.argument_code: item
            for item in argumentation.support_assessments
        }
        conclusion_codes = {
            conclusion.id: conclusion.code
            for conclusion in reasoning.conclusions
        }
        rows: list[TraceabilityRow] = []
        for argument in argumentation.arguments:
            conclusion_code: str | None = None
            if argument.conclusion_id is not None:
                conclusion_code = conclusion_codes.get(
                    argument.conclusion_id,
                    "Conclusión inferida de otra ejecución",
                )
            rows.append(
                TraceabilityRow(
                    argument_code=argument.code,
                    position=argument.position.value,
                    thesis=(
                        f"{argument.thesis_key}="
                        f"{argument.thesis_value.value}"
                    ),
                    conclusion_code=conclusion_code,
                    origin=_argument_origin(argument),
                    fact_codes=argument.fact_codes,
                    evidence_codes=argument.evidence_codes,
                    source_codes=argument.source_codes,
                    rule_codes=argument.rule_codes,
                    assertion_codes=argument.assertion_codes,
                    support_level=support_by_code[argument.code].level.value,
                    support_score=support_by_code[argument.code].score,
                )
            )
        return rows

    @staticmethod
    def _warning_text() -> str:
        return (
            "Este documento organiza información registrada por el usuario y "
            "resultados de un motor determinista. No verifica automáticamente "
            "autenticidad, vigencia, aplicabilidad ni suficiencia jurídica, no "
            "predice decisiones judiciales y no constituye asesoría legal."
        )

    @staticmethod
    def _append_markdown_sources(
        lines: list[str],
        report: IntegralLegalReport,
    ) -> None:
        if not report.issue_source_links:
            lines.append("No hay fuentes vinculadas al problema jurídico.")
            return
        for link in report.issue_source_links:
            lines.append(
                f"- `{link.source_code}` · {link.source_title} · "
                f"{link.orientation.value}. Aplicabilidad: {link.applicability}"
            )

    @staticmethod
    def _append_markdown_arguments(
        lines: list[str],
        report: IntegralLegalReport,
    ) -> None:
        if not report.argumentation.arguments:
            lines.append("No hay argumentos registrados.")
            return
        support_by_code = {
            item.argument_code: item
            for item in report.argumentation.support_assessments
        }
        for argument in report.argumentation.arguments:
            support = support_by_code[argument.code]
            lines.extend(
                [
                    f"### {argument.code} · {argument.title}",
                    "",
                    f"- Posición: {argument.position.value}",
                    f"- Estado: {argument.status.value}",
                    f"- Origen: {_argument_origin(argument)}",
                    (
                        f"- Tesis: `{argument.thesis_key}="
                        f"{argument.thesis_value.value}`"
                    ),
                    (
                        f"- Soporte: {support.level.value} "
                        f"({support.score}/5)"
                    ),
                    "",
                    argument.claim,
                    "",
                ]
            )
            if report.request.include_full_arguments:
                lines.extend([argument.reasoning, ""])
        lines.extend(["### Relaciones", ""])
        labels = {
            argument.id: argument.code
            for argument in report.argumentation.arguments
        }
        for relation in report.argumentation.relations:
            lines.append(
                f"- `{relation.code}`: "
                f"{labels.get(relation.source_argument_id, '?')} "
                f"{relation.relation_type.value.lower()} a "
                f"{labels.get(relation.target_argument_id, '?')}. "
                f"{relation.rationale}"
            )

    @staticmethod
    def _append_markdown_scenarios(
        lines: list[str],
        report: IntegralLegalReport,
    ) -> None:
        lines.extend(
            [
                f"- Base: **{report.base_graph.scenario_name}**",
                f"- Huella base: `{report.base_graph.input_hash}`",
                (
                    "- Objeciones pendientes en base: "
                    f"{report.base_graph.summary['unresolved_objection_count']}"
                ),
            ]
        )
        if report.compared_graph is None:
            lines.append("No se seleccionó un escenario comparado.")
            return
        lines.extend(
            [
                f"- Comparado: **{report.compared_graph.scenario_name}**",
                f"- Huella comparada: `{report.compared_graph.input_hash}`",
                (
                    "- Objeciones pendientes en comparado: "
                    f"{report.compared_graph.summary['unresolved_objection_count']}"
                ),
                "",
                "### Lectura comparativa",
                "",
            ]
        )
        if report.scenario_narrative is not None:
            lines.extend(
                f"- {item}" for item in report.scenario_narrative.statements
            )

    @staticmethod
    def _configure_document(
        document: DocumentObject,
        report: IntegralLegalReport,
    ) -> None:
        section = document.sections[0]
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

        styles = document.styles
        styles["Normal"].font.name = "Arial"
        styles["Normal"].font.size = Pt(10.5)
        for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
            styles[style_name].font.name = "Arial"

        header = section.header.paragraphs[0]
        header.text = "IUS-Razón · Informe jurídico integral"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header.runs[0].font.size = Pt(8)

        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("IUS-Razón · Powered by Sképsis Apps · Página ")
        _add_page_number(footer)
        footer.add_run(f" · {report.case.confidentiality.value}")
        for run in footer.runs:
            run.font.size = Pt(8)

    @staticmethod
    def _add_cover(
        document: DocumentObject,
        report: IntegralLegalReport,
    ) -> None:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(24)
        run = paragraph.add_run("IUS-RAZÓN")
        run.bold = True
        run.font.size = Pt(20)

        title = document.add_heading(report.request.title, 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.add_paragraph("")
        details = document.add_paragraph()
        details.alignment = WD_ALIGN_PARAGRAPH.CENTER
        details.add_run(f"Expediente\n{report.case.title}\n\n").bold = True
        details.add_run(
            f"Problema {report.issue.code}\n{report.issue.question}\n\n"
            f"Jurisdicción: {report.case.jurisdiction}\n"
            f"Fecha de instantánea: {report.generated_at.isoformat()}\n"
            f"Versión: {report.report_version}\n"
            f"Huella SHA-256:\n{report.input_hash}"
        )
        warning = document.add_paragraph(report.warning)
        warning.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    @staticmethod
    def _add_contents(
        document: DocumentObject,
        report: IntegralLegalReport,
    ) -> None:
        document.add_heading("Índice", level=1)
        for item in report.section_index:
            document.add_paragraph(item)

    @staticmethod
    def _add_identification(
        document: DocumentObject,
        report: IntegralLegalReport,
    ) -> None:
        document.add_heading("1. Identificación y alcance", level=1)
        document.add_paragraph(report.request.purpose)
        table = document.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        headers = table.rows[0].cells
        headers[0].text = "Campo"
        headers[1].text = "Valor"
        _mark_table_header(table.rows[0])
        rows = [
            ("Expediente", report.case.title),
            ("Materia", report.case.matter),
            ("Jurisdicción", report.case.jurisdiction),
            ("Problema", f"{report.issue.code} · {report.issue.title}"),
            ("Pregunta", report.issue.question),
            ("Estado", report.case.status.value),
            ("Confidencialidad", report.case.confidentiality.value),
        ]
        for label, value in rows:
            cells = table.add_row().cells
            cells[0].text = label
            cells[1].text = value
            cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        document.add_heading("Partes registradas", level=2)
        if not report.parties:
            document.add_paragraph("No hay partes registradas.")
        for party in report.parties:
            document.add_paragraph(
                f"{party.name_alias} · {party.legal_role.value} · "
                f"{party.party_type.value}.",
                style="List Bullet",
            )

    @staticmethod
    def _add_executive_summary(
        document: DocumentObject,
        report: IntegralLegalReport,
    ) -> None:
        document.add_heading("2. Resumen ejecutivo", level=1)
        document.add_paragraph(report.executive_summary)

    @staticmethod
    def _add_methodology(
        document: DocumentObject,
        report: IntegralLegalReport,
    ) -> None:
        document.add_heading("3. Metodología y advertencias", level=1)
        for item in report.methodology:
            document.add_paragraph(item, style="List Bullet")
        paragraph = document.add_paragraph(report.warning)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    @staticmethod
    def _add_background(
        document: DocumentObject,
        report: IntegralLegalReport,
    ) -> None:
        document.add_heading("4. Antecedentes del expediente", level=1)
        document.add_paragraph(report.case.description)
        document.add_paragraph(
            f"Objetivo registrado: {report.case.objective or 'No registrado'}"
        )
        document.add_paragraph(
            f"Rol del usuario: {report.case.user_role or 'No registrado'}"
        )

    @staticmethod
    def _add_facts_and_evidence(
        document: DocumentObject,
        report: IntegralLegalReport,
    ) -> None:
        document.add_heading("5. Hechos y pruebas", level=1)
        document.add_heading("Hechos", level=2)
        if not report.facts:
            document.add_paragraph("No hay hechos registrados.")
        for fact in report.facts:
            document.add_paragraph(
                f"{fact.code} · {fact.status.value} · {fact.description}",
                style="List Bullet",
            )
        document.add_heading("Pruebas", level=2)
        if not report.evidence:
            document.add_paragraph("No hay pruebas registradas.")
        for item in report.evidence:
            document.add_paragraph(
                f"{item.code} · {item.evidence_type.value} · "
                f"{item.evaluation_status.value} · {item.description}",
                style="List Bullet",
            )
        document.add_heading("Vínculos hecho-prueba", level=2)
        if not report.fact_evidence_links:
            document.add_paragraph("No hay vínculos hecho-prueba registrados.")
        for link in report.fact_evidence_links:
            document.add_paragraph(
                f"{link.fact_code} ↔ {link.evidence_code} · "
                f"{_ensure_sentence(link.purpose or 'Sin propósito registrado')}",
                style="List Bullet",
            )

    @staticmethod
    def _add_sources(
        document: DocumentObject,
        report: IntegralLegalReport,
    ) -> None:
        document.add_heading("6. Fuentes jurídicas vinculadas", level=1)
        if not report.issue_source_links:
            document.add_paragraph(
                "No hay fuentes vinculadas al problema jurídico."
            )
            return
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        headers = table.rows[0].cells
        headers[0].text = "Código"
        headers[1].text = "Fuente"
        headers[2].text = "Orientación"
        headers[3].text = "Aplicabilidad"
        _mark_table_header(table.rows[0])
        for link in report.issue_source_links:
            cells = table.add_row().cells
            cells[0].text = link.source_code
            cells[1].text = link.source_title
            cells[2].text = link.orientation.value
            cells[3].text = link.applicability
        if report.request.include_source_details:
            LegalReportService._add_source_details(document, report)

    @staticmethod
    def _add_source_details(
        document: DocumentObject,
        report: IntegralLegalReport,
    ) -> None:
        document.add_heading("Detalle de fuentes", level=2)
        for norm in report.norms:
            document.add_heading(
                f"{norm.code} · {norm.instrument} · {norm.article}",
                level=3,
            )
            document.add_paragraph(
                f"Jerarquía: {norm.hierarchy.value}. "
                f"Jurisdicción: {norm.jurisdiction}. "
                f"Texto registrado: {norm.text}"
            )
        for precedent in report.jurisprudence:
            document.add_heading(
                f"{precedent.code} · {precedent.court} · {precedent.identifier}",
                level=3,
            )
            document.add_paragraph(
                f"Autoridad: {precedent.authority.value}. "
                f"Criterio: {precedent.criterion}"
            )
        for doctrine in report.doctrine:
            document.add_heading(
                f"{doctrine.code} · {doctrine.author} · {doctrine.work_title}",
                level=3,
            )
            document.add_paragraph(
                f"Concepto: {doctrine.concept}. "
                f"Posición: {doctrine.position_summary}"
            )

    @staticmethod
    def _add_reasoning(
        document: DocumentObject,
        report: IntegralLegalReport,
    ) -> None:
        document.add_heading("7. Inferencia y trazabilidad", level=1)
        document.add_paragraph(
            f"Ejecución: {report.reasoning.run.id}\n"
            f"Motor: {report.reasoning.run.engine_version}\n"
            f"Huella: {report.reasoning.run.input_hash}\n"
            f"Estado: {report.reasoning.run.status.value}"
        )
        document.add_heading("Conclusiones inferidas", level=2)
        if not report.reasoning.conclusions:
            document.add_paragraph(
                "No se derivaron conclusiones en la ejecución seleccionada."
            )
        for conclusion in report.reasoning.conclusions:
            document.add_heading(
                f"{conclusion.code} · {conclusion.predicate_key}",
                level=3,
            )
            document.add_paragraph(conclusion.statement)
            document.add_paragraph(
                f"Valor: {conclusion.value.value}\n"
                f"Estado: {conclusion.status.value}\n"
                f"Soporte: {conclusion.support_level.value}\n"
                f"Reglas: {_codes(conclusion.rule_codes)}\n"
                "Premisas: "
                f"{_codes(conclusion.supporting_assertion_codes)}\n"
                f"Fuentes: {_codes(conclusion.source_codes)}"
            )
        document.add_heading("Traza de reglas", level=2)
        if not report.reasoning.traces:
            document.add_paragraph("No hay trazas registradas.")
        for trace in report.reasoning.traces:
            document.add_paragraph(
                f"{trace.sequence}. {trace.rule_code} · "
                f"{trace.outcome.value} · "
                f"{json.dumps(trace.detail, ensure_ascii=False, sort_keys=True)}",
                style="List Bullet",
            )

    @staticmethod
    def _add_argumentation(
        document: DocumentObject,
        report: IntegralLegalReport,
    ) -> None:
        document.add_heading("8. Argumentación jurídica", level=1)
        support_by_code = {
            item.argument_code: item
            for item in report.argumentation.support_assessments
        }
        if not report.argumentation.arguments:
            document.add_paragraph("No hay argumentos registrados.")
        for argument in report.argumentation.arguments:
            support = support_by_code[argument.code]
            document.add_heading(
                f"{argument.code} · {argument.title}",
                level=2,
            )
            document.add_paragraph(
                f"Posición: {argument.position.value}\n"
                f"Estado: {argument.status.value}\n"
                f"Origen: {_argument_origin(argument)}\n"
                f"Tesis: {argument.thesis_key}="
                f"{argument.thesis_value.value}\n"
                f"Soporte: {support.level.value} ({support.score}/5)"
            )
            document.add_paragraph(argument.claim)
            if report.request.include_full_arguments:
                document.add_paragraph(argument.reasoning)
        document.add_heading("Relaciones argumentales", level=2)
        labels = {
            argument.id: argument.code
            for argument in report.argumentation.arguments
        }
        if not report.argumentation.relations:
            document.add_paragraph("No hay relaciones registradas.")
        for relation in report.argumentation.relations:
            document.add_paragraph(
                f"{relation.code}: "
                f"{labels.get(relation.source_argument_id, '?')} "
                f"{relation.relation_type.value.lower()} a "
                f"{labels.get(relation.target_argument_id, '?')}. "
                f"{relation.rationale}",
                style="List Bullet",
            )

    @staticmethod
    def _add_scenarios(
        document: DocumentObject,
        report: IntegralLegalReport,
    ) -> None:
        document.add_heading("9. Escenarios alternativos", level=1)
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        headers = table.rows[0].cells
        headers[0].text = "Escenario"
        headers[1].text = "Argumentos"
        headers[2].text = "Relaciones"
        headers[3].text = "Objeciones pendientes"
        _mark_table_header(table.rows[0])
        graphs = [report.base_graph]
        if report.compared_graph is not None:
            graphs.append(report.compared_graph)
        for graph in graphs:
            cells = table.add_row().cells
            cells[0].text = f"{graph.scenario_code or 'COMPLETO'} · {graph.scenario_name}"
            cells[1].text = str(graph.summary["node_count"])
            cells[2].text = str(graph.summary["edge_count"])
            cells[3].text = str(
                graph.summary["unresolved_objection_count"]
            )
        if report.scenario_narrative is not None:
            document.add_heading("Lectura comparativa", level=2)
            for item in report.scenario_narrative.statements:
                document.add_paragraph(item, style="List Bullet")
        else:
            document.add_paragraph(
                "No se seleccionó un escenario comparado."
            )

    @staticmethod
    def _add_findings(
        document: DocumentObject,
        report: IntegralLegalReport,
    ) -> None:
        document.add_heading("10. Hallazgos y conclusiones", level=1)
        for item in report.findings:
            document.add_paragraph(item, style="List Bullet")
        if report.request.analyst_conclusions:
            document.add_heading(
                "Conclusiones registradas por el analista",
                level=2,
            )
            for item in report.request.analyst_conclusions:
                document.add_paragraph(item, style="List Bullet")

    @staticmethod
    def _add_limitations(
        document: DocumentObject,
        report: IntegralLegalReport,
    ) -> None:
        document.add_heading(
            "11. Limitaciones e información faltante",
            level=1,
        )
        for item in report.limitations:
            document.add_paragraph(item, style="List Bullet")

    @staticmethod
    def _add_recommendations(
        document: DocumentObject,
        report: IntegralLegalReport,
    ) -> None:
        document.add_heading("12. Recomendaciones de revisión", level=1)
        for item in report.recommendations:
            document.add_paragraph(item, style="List Bullet")

    @staticmethod
    def _add_traceability_appendix(
        document: DocumentObject,
        report: IntegralLegalReport,
    ) -> None:
        document.add_section(WD_SECTION.NEW_PAGE)
        document.add_heading("Anexo A. Matriz de trazabilidad", level=1)
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        headers = table.rows[0].cells
        headers[0].text = "Argumento"
        headers[1].text = "Tesis"
        headers[2].text = "Soporte"
        headers[3].text = "Trazabilidad"
        _mark_table_header(table.rows[0])
        for row in report.traceability_rows:
            cells = table.add_row().cells
            cells[0].text = f"{row.argument_code}\n{row.position}"
            cells[1].text = row.thesis
            cells[2].text = f"{row.support_level} ({row.support_score}/5)"
            cells[3].text = (
                f"Origen: {row.origin}\n"
                f"Conclusión: "
                f"{row.conclusion_code or 'No aplica (argumento manual)'}\n"
                f"Hechos: {_codes(row.fact_codes)}\n"
                f"Pruebas: {_codes(row.evidence_codes)}\n"
                f"Fuentes: {_codes(row.source_codes)}\n"
                f"Reglas: {_codes(row.rule_codes)}\n"
                f"Premisas: {_codes(row.assertion_codes)}"
            )

    @staticmethod
    def _add_reproducibility_appendix(
        document: DocumentObject,
        report: IntegralLegalReport,
    ) -> None:
        document.add_heading(
            "Anexo B. Huellas y reproducibilidad",
            level=1,
        )
        document.add_paragraph(
            f"Informe: {report.input_hash}\n"
            f"Inferencia: {report.reasoning.run.input_hash}\n"
            f"Grafo base: {report.base_graph.input_hash}"
        )
        if report.compared_graph is not None:
            document.add_paragraph(
                f"Grafo comparado: {report.compared_graph.input_hash}"
            )
        document.add_heading("Advertencia final", level=2)
        document.add_paragraph(report.warning)
